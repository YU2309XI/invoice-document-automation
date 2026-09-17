from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)


DEFAULT_OUTPUT_FILE = Path("output/invoice_report.xlsx")


def style_worksheet(worksheet):
    """Style invoice data worksheets."""

    columns = {}

    for cell in worksheet[1]:
        columns[cell.value] = cell.column

    # Freeze header
    worksheet.freeze_panes = "A2"

    # Enable filter
    worksheet.auto_filter.ref = worksheet.dimensions

    # Header style
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E1F2",
        )
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 24

    # Data rows
    status_column = columns.get("status")

    for row in range(2, worksheet.max_row + 1):
        if status_column:
            status_cell = worksheet.cell(
                row=row,
                column=status_column,
            )

            if status_cell.value == "OK":
                status_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="E2F0D9",
                )

                status_cell.font = Font(
                    color="375623",
                    bold=True,
                )

            elif status_cell.value == "NEEDS REVIEW":
                status_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FCE4D6",
                )

                status_cell.font = Font(
                    color="C00000",
                    bold=True,
                )

        for cell in worksheet[row]:
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="top",
            )

    # Money formatting
    for field in ["subtotal", "tax", "total"]:
        if field not in columns:
            continue

        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(
                row=row,
                column=columns[field],
            ).number_format = "#,##0.00"

    # Column widths
    widths = {
        "source_file": 20,
        "extraction_method": 18,
        "vendor": 30,
        "invoice_number": 20,
        "invoice_date": 16,
        "due_date": 16,
        "customer": 30,
        "currency": 12,
        "subtotal": 14,
        "tax": 12,
        "total": 14,
        "status": 18,
        "issues": 48,
    }

    for field, width in widths.items():
        if field not in columns:
            continue

        column_number = columns[field]

        column_letter = worksheet.cell(
            row=1,
            column=column_number,
        ).column_letter

        worksheet.column_dimensions[
            column_letter
        ].width = width

    # Wrap issue descriptions
    if "issues" in columns:
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(
                row=row,
                column=columns["issues"],
            ).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )


def style_summary(workbook):
    """Style the Summary dashboard."""

    if "Summary" not in workbook.sheetnames:
        return

    worksheet = workbook["Summary"]

    # Cleaner dashboard appearance
    worksheet.sheet_view.showGridLines = False

    # Hide pandas Metric / Value header
    worksheet.row_dimensions[1].hidden = True

    # Main title
    worksheet.merge_cells("A2:F2")

    title = worksheet["A2"]
    title.value = "INVOICE PROCESSING SUMMARY"

    title.font = Font(
        size=20,
        bold=True,
        color="FFFFFF",
    )

    title.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    title.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[2].height = 38

    # Subtitle
    worksheet.merge_cells("A3:F3")

    subtitle = worksheet["A3"]

    subtitle.value = (
        "Automated PDF • OCR • Validation • Excel Reporting"
    )

    subtitle.font = Font(
        size=11,
        italic=True,
        color="666666",
    )

    subtitle.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[3].height = 24

    # Section headings
    section_rows = [4, 9, 13]

    for row in section_rows:
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=6,
        )

        cell = worksheet.cell(
            row=row,
            column=1,
        )

        cell.font = Font(
            size=12,
            bold=True,
            color="1F4E78",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        cell.alignment = Alignment(
            vertical="center",
        )

        worksheet.row_dimensions[row].height = 24

    # Metric rows
    metric_rows = [5, 6, 7, 10, 11, 14, 15]

    for row in metric_rows:
        worksheet[f"A{row}"].font = Font(
            bold=True,
            color="404040",
        )

        worksheet[f"B{row}"].font = Font(
            size=14,
            bold=True,
            color="1F4E78",
        )

        worksheet[f"B{row}"].alignment = Alignment(
            horizontal="center",
        )

    # Successful count
    worksheet["B6"].fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    worksheet["B6"].font = Font(
        size=14,
        bold=True,
        color="375623",
    )

    # Needs Review count
    worksheet["B7"].fill = PatternFill(
        fill_type="solid",
        fgColor="FCE4D6",
    )

    worksheet["B7"].font = Font(
        size=14,
        bold=True,
        color="C00000",
    )

    # Currency formatting
    worksheet["B14"].number_format = '€#,##0.00'
    worksheet["B15"].number_format = '$#,##0.00'

    # Dashboard width
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 20

    for column in ["C", "D", "E", "F"]:
        worksheet.column_dimensions[column].width = 12


def style_report(output_file=None):
    output_file = Path(output_file) if output_file else DEFAULT_OUTPUT_FILE

    workbook = load_workbook(output_file)

    # Style detailed data sheets
    for sheet_name in [
        "Invoices",
        "Needs Review",
    ]:
        if sheet_name in workbook.sheetnames:
            style_worksheet(
                workbook[sheet_name]
            )

    # Style dashboard
    style_summary(workbook)

    workbook.save(output_file)

    


if __name__ == "__main__":
    style_report()
